import csv, io
from datetime import date, datetime
from decimal import Decimal
from xml.etree.ElementTree import Element, SubElement, tostring
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def _normalize(value):
    if isinstance(value, (datetime, date)): return value.isoformat()
    if isinstance(value, Decimal): return float(value)
    if isinstance(value, (int, float, str, bool)) or value is None: return value
    return str(value)

def to_csv(rows: list[dict], fields: list[str]) -> bytes:
    buf = io.StringIO(); w = csv.DictWriter(buf, fieldnames=fields, extrasaction="ignore", delimiter=";")
    w.writeheader(); w.writerows(rows); return buf.getvalue().encode("utf-8-sig")

def to_xlsx(rows: list[dict], fields: list[str], sheet_title: str = "data") -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = sheet_title[:31] or "data"; ws.append(fields)
    for r in rows: ws.append([r.get(f) for f in fields])
    for i, f in enumerate(fields, start=1):
        max_len = max([len(str(f))] + [len(str(r.get(f))) for r in rows]) if rows else len(f)
        ws.column_dimensions[get_column_letter(i)].width = min(50, max_len + 2)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

def to_xml(rows: list[dict], root_name: str = "rows", item_name: str = "row") -> bytes:
    root = Element(root_name)
    for r in rows:
        item = SubElement(root, item_name)
        for k, v in r.items():
            child = SubElement(item, k); child.text = "" if v is None else str(v)
    return tostring(root, encoding="utf-8", xml_declaration=True)

FORMATS = {"csv": ("text/csv; charset=utf-8", "csv"), "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"), "xml": ("application/xml; charset=utf-8", "xml")}

def export(rows: list[dict], fields: list[str], fmt: str, filename_base: str) -> tuple[bytes, str, str]:
    fmt = fmt.lower()
    if fmt not in FORMATS: raise ValueError(f"Неподдерживаемый формат: {fmt}")
    mime, ext = FORMATS[fmt]
    if fmt == "csv": data = to_csv(rows, fields)
    elif fmt == "xlsx": data = to_xlsx(rows, fields, sheet_title=filename_base)
    else: data = to_xml(rows, root_name=filename_base, item_name="item")
    return data, mime, f"{filename_base}.{ext}"