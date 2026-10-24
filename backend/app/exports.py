"""Tabular exports: CSV for Excel, XLSX and XML.

Text cells that a spreadsheet would run as a formula are neutralised: CSV gets a
leading apostrophe, XLSX cells are stored as plain text.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from xml.etree.ElementTree import Element, SubElement, tostring

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
FORMATS = {
    "csv": ("text/csv; charset=utf-8", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "xml": ("application/xml; charset=utf-8", "xml"),
}


def _plain(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def to_csv(rows: list[dict[str, Any]], fields: list[str]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(fields)
    for row in rows:
        cells = []
        for field in fields:
            value = _plain(row.get(field))
            if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
                value = "'" + value
            cells.append("" if value is None else value)
        writer.writerow(cells)
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(rows: list[dict[str, Any]], fields: list[str], title: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "data"
    sheet.append(fields)
    for row_number, row in enumerate(rows, start=2):
        for column, field in enumerate(fields, start=1):
            value = _plain(row.get(field))
            cell = sheet.cell(row=row_number, column=column, value=value)
            if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
                cell.data_type = "s"
    for column, field in enumerate(fields, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = min(40, max(len(field) + 2, 12))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_xml(rows: list[dict[str, Any]], fields: list[str], root_name: str) -> bytes:
    root = Element(root_name)
    for row in rows:
        item = SubElement(root, "item")
        for field in fields:
            child = SubElement(item, field)
            value = _plain(row.get(field))
            child.text = "" if value is None else str(value)
    return tostring(root, encoding="utf-8", xml_declaration=True)


def export(rows: list[dict[str, Any]], fields: list[str], fmt: str, base_name: str) -> tuple[bytes, str, str]:
    media_type, extension = FORMATS[fmt]
    if fmt == "csv":
        data = to_csv(rows, fields)
    elif fmt == "xlsx":
        data = to_xlsx(rows, fields, base_name)
    else:
        data = to_xml(rows, fields, base_name.replace("-", "_"))
    return data, media_type, f"{base_name}.{extension}"
