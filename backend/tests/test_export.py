import csv
import io
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

EXPORT = "/api/v1/export/{entity}"
ADMIN_PASSWORD = "local-demo-change-this-password"


def sign_in(client) -> None:
    assert client.post("/api/v1/auth/login", json={"username": "demo", "password": ADMIN_PASSWORD}).status_code == 200


def test_customers_csv_has_a_header_and_one_row_per_customer(open_app):
    with open_app() as (client, _):
        sign_in(client)
        response = client.get(EXPORT.format(entity="customers"), params={"format": "csv"})
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert response.headers["content-disposition"].startswith("attachment")
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
        assert rows[0][:2] == ["id", "email"]
        assert len(rows) - 1 == 10


def test_xlsx_export_opens_as_a_workbook(open_app):
    with open_app() as (client, _):
        sign_in(client)
        response = client.get(EXPORT.format(entity="services"), params={"format": "xlsx"})
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet["A1"].value == "id"
        assert sheet.max_row - 1 == 30


def test_xml_export_is_well_formed(open_app):
    with open_app() as (client, _):
        sign_in(client)
        response = client.get(EXPORT.format(entity="transactions"), params={"format": "xml"})
        assert response.status_code == 200
        assert len(list(ET.fromstring(response.content))) == 100


def test_risk_queue_export_carries_the_recommended_action(open_app):
    with open_app() as (client, _):
        sign_in(client)
        response = client.get(EXPORT.format(entity="risk"), params={"format": "csv"})
        header = response.content.decode("utf-8-sig").splitlines()[0].split(";")
        assert {"risk_score", "risk_level", "recommended_action_text", "near_term_loss_minor"} <= set(header)


def test_sensitive_tables_are_not_exported(open_app):
    with open_app() as (client, _):
        sign_in(client)
        for entity in ("sessions", "users", "login_attempts", "app_state", "password_resets"):
            assert client.get(EXPORT.format(entity=entity)).status_code == 404


def test_unknown_format_is_rejected(open_app):
    with open_app() as (client, _):
        sign_in(client)
        assert client.get(EXPORT.format(entity="customers"), params={"format": "pdf"}).status_code == 422


def test_export_requires_a_session(open_app):
    with open_app() as (client, _):
        assert client.get(EXPORT.format(entity="customers")).status_code == 401


def test_formula_like_values_are_neutralised(open_app):
    with open_app() as (client, _):
        sign_in(client)
        payload = {"customers": [{
            "id": "9", "email": '=HYPERLINK("http://example.test")@example.test', "balance": 0,
            "services": [], "transactions": [],
        }]}
        assert client.put("/api/v1/source/snapshot", json=payload).status_code == 200
        csv_text = client.get(EXPORT.format(entity="customers"), params={"format": "csv"}).content.decode("utf-8-sig")
        assert "'=HYPERLINK" in csv_text
        sheet = load_workbook(io.BytesIO(client.get(EXPORT.format(entity="customers"), params={"format": "xlsx"}).content)).active
        assert sheet["B2"].data_type == "s"
        assert sheet["B2"].value.startswith("=HYPERLINK")
